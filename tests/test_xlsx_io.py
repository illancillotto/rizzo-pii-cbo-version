"""Synthetic-only OOXML flow tests. The fake detector does NOT test model quality."""
import copy
import datetime
import io
import json
import sys
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src/app"))
try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Font
    from openpyxl.workbook.defined_name import DefinedName
    from xlsx_io import Limits, Package, WorkbookJob, XlsxError, restore
    from pii_session import PiiSession
    from detectors import detect_regex
    AVAILABLE = True
except ImportError:
    AVAILABLE = False

NAME = "Alfa Esempio"
EMAIL = "sintetico@example.invalid"


def fake_batch(texts):
    out = []
    for text in texts:
        entities = detect_regex(text)
        start = 0
        while (start := text.find(NAME, start)) >= 0:
            entities.append({"start": start, "end": start + len(NAME), "label": "FULLNAME"})
            start += len(NAME)
        kept = []
        for ent in sorted(entities, key=lambda e: (e["start"], -(e["end"] - e["start"]))):
            if not kept or ent["start"] >= kept[-1]["end"]:
                kept.append(ent)
        out.append(kept)
    return out


def workbook(values=None):
    book = Workbook()
    for coordinate, value in (values or {"A1": "Intestazione", "A2": EMAIL}).items():
        book.active[coordinate] = value
    return book


def save(book):
    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()


def mutate(data, updates):
    stream = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as source, zipfile.ZipFile(stream, "w", zipfile.ZIP_DEFLATED) as target:
        for part in source.namelist():
            value = source.read(part)
            if part in updates:
                value = updates[part](value) if callable(updates[part]) else updates[part]
            target.writestr(part, value)
        for part, value in updates.items():
            if part not in source.namelist():
                target.writestr(part, value)
    return stream.getvalue()


@unittest.skipUnless(AVAILABLE, "Installare requirements-xlsx-test.txt (job XLSX obbligatorio in CI)")
class XlsxFlow(unittest.TestCase):
    def job(self, book=None, config=None):
        job = WorkbookJob(save(book or workbook()))
        job.analyze(fake_batch, config or {})
        return job

    def test_same_value_across_sheets_and_session_isolation(self):
        book = workbook({"A1": NAME})
        book.create_sheet("Secondo")["D2"] = NAME
        job = self.job(book)
        self.assertEqual(len({c["replacement"] for c in job.changes}), 1)
        other = self.job(book)
        self.assertNotEqual(job.changes[0]["replacement"], other.changes[0]["replacement"])
        result = load_workbook(io.BytesIO(job.export(fake_batch)))
        self.assertEqual(result.worksheets[0]["A1"].value, result.worksheets[1]["D2"].value)

    def test_exact_identity_category_and_spelling(self):
        session = PiiSession(["already here"])
        values = [session.placeholder("ID", value) for value in ["Alfa", "alfa", " Alfa", "Alfa ", "Alfa\n"]]
        self.assertEqual(len(set(values)), 5)
        self.assertNotEqual(session.placeholder("ID", "Alfa"), session.placeholder("FULLNAME", "Alfa"))
        self.assertNotEqual(session.placeholder("ID", "42", "n"), session.placeholder("ID", "42"))
        self.assertEqual(session.placeholder("ID", "Alfa"), values[0])

    def test_existing_placeholder_and_namespace_collision(self):
        existing = "[[RZ_aaaaaaaaaaaaaaaa_ID_1]]"
        with patch("pii_session.secrets.token_hex", side_effect=["aaaaaaaaaaaaaaaa", "bbbbbbbbbbbbbbbb"]):
            session = PiiSession([existing])
        self.assertIn("bbbbbbbbbbbbbbbb", session.placeholder("ID", "new"))
        job = self.job(workbook({"A1": existing, "A2": EMAIL}))
        restored = load_workbook(io.BytesIO(restore(job.export(fake_batch), job.dictionary)))
        self.assertEqual(restored.active["A1"].value, existing)

    def test_regex_identifiers_and_free_text(self):
        values = {"A1": NAME, "A2": "RSSMRA85H12F205Y", "A3": "12345678903",
                  "A4": "IT60X0542811101000000123456", "A5": "Riga\n" + EMAIL + "\nΩ"}
        job = self.job(workbook(values))
        categories = {c["category"] for c in job.changes}
        self.assertTrue({"FULLNAME", "CF", "PIVA", "IBAN", "EMAIL"} <= categories)
        exported = job.export(fake_batch)
        restored = load_workbook(io.BytesIO(restore(exported, job.dictionary)))
        for coordinate, value in values.items():
            self.assertEqual(restored.active[coordinate].value, value)

    def test_forced_numeric_identifiers_headers_and_exclusion_precedence(self):
        book = workbook({"A1": "Codice", "A2": 12345678, "A3": "000123", "B2": 19.5})
        config = {"exclude_tags": ["ID", "AMOUNT", "DATE"], "columns": [
            {"sheet": 0, "column": 1, "header_row": 1, "start_row": 2, "category": "ID"}]}
        job = self.job(book, config)
        self.assertEqual(len(job.changes), 2)
        self.assertTrue(all(c["conflict"] for c in job.changes))
        with self.assertRaises(XlsxError):
            job.correct({"id": 0, "enabled": False})
        output = job.export(fake_batch)
        masked = load_workbook(io.BytesIO(output))
        self.assertEqual(masked.active["A1"].value, "Codice")
        self.assertEqual(masked.active["B2"].value, 19.5)
        recovered = load_workbook(io.BytesIO(restore(output, job.dictionary)))
        self.assertEqual(recovered.active["A2"].value, 12345678)
        self.assertEqual(recovered.active["A2"].data_type, "n")
        self.assertEqual(recovered.active["A3"].value, "000123")

    def test_hidden_merged_unicode_empty_styles_dimensions(self):
        book = workbook({"A1": NAME, "C3": "Ω\n" + EMAIL})
        book.active.merge_cells("A1:B1")
        book.active.row_dimensions[3].hidden = True
        book.active.column_dimensions["C"].hidden = True
        book.active.column_dimensions["C"].width = 31
        book.active["C3"].font = Font(bold=True, color="FF0000")
        book.create_sheet("Nascosto")["B2"] = EMAIL
        book.worksheets[1].sheet_state = "hidden"
        book.create_sheet("MoltoNascosto")["B2"] = NAME
        book.worksheets[2].sheet_state = "veryHidden"
        job = self.job(book)
        out = load_workbook(io.BytesIO(job.export(fake_batch)))
        self.assertEqual(out.worksheets[2].sheet_state, "veryHidden")
        self.assertTrue(out.active.row_dimensions[3].hidden)
        self.assertTrue(out.active.column_dimensions["C"].hidden)
        self.assertEqual(out.active.column_dimensions["C"].width, 31)
        self.assertTrue(out.active["C3"].font.bold)
        self.assertIn("A1:B1", out.active.merged_cells)
        self.assertIsNone(out.active["B1"].value)

    def test_boolean_date_error_and_number_restore(self):
        book = workbook({"A1": "Tipo", "A2": True, "A3": datetime.datetime(2025, 1, 2),
                         "A4": "#N/A", "A5": 78123})
        job = self.job(book, {"columns": [{"sheet": 0, "column": 1}], "exclude_tags": ["DATE"]})
        out = job.export(fake_batch)
        restored = load_workbook(io.BytesIO(restore(out, job.dictionary)))
        for row in range(2, 6):
            self.assertEqual(restored.active.cell(row, 1).value, book.active.cell(row, 1).value)
            self.assertEqual(restored.active.cell(row, 1).data_type, book.active.cell(row, 1).data_type)

    def test_metadata_comments_headers_and_restore(self):
        book = workbook()
        book.properties.creator = NAME
        book.properties.description = EMAIL
        book.active["C1"].comment = Comment("Nota " + EMAIL, NAME)
        book.active.oddHeader.center.text = NAME
        book.active.oddFooter.left.text = EMAIL
        job = self.job(book)
        out = job.export(fake_batch)
        with zipfile.ZipFile(io.BytesIO(out)) as archive:
            for part in archive.namelist():
                self.assertNotIn(EMAIL.encode(), archive.read(part))
                self.assertNotIn(NAME.encode(), archive.read(part))
        restored = load_workbook(io.BytesIO(restore(out, job.dictionary)))
        self.assertEqual(restored.properties.creator, NAME)
        self.assertEqual(restored.active["C1"].comment.author, NAME)
        self.assertEqual(restored.active.oddHeader.center.text, NAME)

    def test_hyperlink_and_personal_names_blocked(self):
        book = workbook()
        book.active["B1"].hyperlink = "https://example.invalid/"
        with self.assertRaisesRegex(XlsxError, "hyperlink"):
            WorkbookJob(save(book))
        book = workbook()
        book.active.title = NAME
        with self.assertRaisesRegex(XlsxError, "nomi"):
            self.job(book)

    def test_defined_names_static_only(self):
        book = workbook()
        book.defined_names.add(DefinedName("Area", attr_text="'Sheet'!$A$1:$A$2"))
        self.job(book).export(fake_batch)
        book.defined_names.add(DefinedName("Literal", attr_text='"' + EMAIL + '"'))
        with self.assertRaisesRegex(XlsxError, "Nomi definiti"):
            WorkbookJob(save(book))

    def test_formulas_supported_cache_removed_and_dependencies_blocked(self):
        book = workbook({"A1": 7, "A2": 8, "B1": "=SUM(A1:A2)", "C1": EMAIL})
        job = self.job(book)
        out = job.export(fake_batch)
        self.assertEqual(load_workbook(io.BytesIO(out)).active["B1"].value, "=SUM(A1:A2)")
        self.assertIsNone(load_workbook(io.BytesIO(out), data_only=True).active["B1"].value)
        job = self.job(book, {"columns": [{"sheet": 0, "column": 1, "start_row": 1, "header_row": 0}]})
        with self.assertRaisesRegex(XlsxError, "dipende"):
            job.export(fake_batch)

    def test_unsafe_formulas_blocked(self):
        for formula in ['="' + EMAIL + '"', "='Other'!A1", "=[other.xlsx]Sheet!A1",
                        '=HYPERLINK("https://example.invalid")', "=NOW()", "=DefinedName", "=A:A"]:
            with self.subTest(formula=formula), self.assertRaises(XlsxError):
                WorkbookJob(save(workbook({"A1": formula})))

    def test_complex_parts_and_corruption(self):
        original = save(workbook())
        for part in ["xl/vbaProject.bin", "xl/media/image1.png", "xl/pivotCache/pivotCacheDefinition1.xml",
                     "xl/connections.xml", "xl/embeddings/oleObject1.bin", "docProps/custom.xml"]:
            with self.subTest(part=part), self.assertRaises(XlsxError):
                WorkbookJob(mutate(original, {part: b"synthetic"}))
        for raw in [b"not a zip", b"PK\x03\x04bad", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1fake"]:
            with self.assertRaises(XlsxError):
                WorkbookJob(raw)
        for extension in ["xls", "xlsm", "xlsb"]:
            with self.assertRaises(XlsxError):
                WorkbookJob(original, "document." + extension)

    def test_zip_cell_sheet_character_and_time_limits(self):
        raw = save(workbook())
        for limits in [Limits(compressed=1), Limits(expanded=1), Limits(parts=1),
                       Limits(cells=1), Limits(cell_length=4), Limits(characters=1)]:
            with self.subTest(limits=limits), self.assertRaises(XlsxError):
                WorkbookJob(raw, limits=limits)
        book = workbook(); book.create_sheet("Secondo")
        with self.assertRaises(XlsxError):
            WorkbookJob(save(book), limits=Limits(sheets=1))
        job = self.job()
        with patch("xlsx_io.time.monotonic", side_effect=[0, 10000]), self.assertRaises(XlsxError):
            job.analyze(fake_batch, {})

    def test_no_mapping_has_no_original_in_preview_or_dictionary(self):
        job = self.job(config={"include_mapping": False})
        self.assertNotIn(EMAIL, json.dumps(job.preview()))
        self.assertEqual(job.session.mapping, {})
        job.export(fake_batch)
        self.assertIsNone(job.dictionary)

    def test_review_occurrence_associated_and_add(self):
        job = self.job(workbook({"A1": EMAIL, "A2": EMAIL, "B1": "interno-123"}))
        job.correct({"id": 0, "category": "ID", "scope": "associated"})
        self.assertTrue(all(c["category"] == "ID" for c in job.changes))
        job.correct({"id": 0, "enabled": False})
        self.assertTrue(job.changes[1]["enabled"])
        job.add({"sheet": 0, "cell": "B1", "category": "ACCOUNT"})
        job.export(fake_batch)
        job.correct({"id": 0, "enabled": True})
        self.assertIsNone(job.output)
        out = job.export(fake_batch)
        self.assertNotEqual(load_workbook(io.BytesIO(out)).active["B1"].value, "interno-123")

    def test_pagination_and_validation(self):
        job = self.job(workbook({f"A{n}": EMAIL for n in range(1, 121)}))
        self.assertEqual(len(job.preview(1, 50)["items"]), 50)
        self.assertEqual(job.preview(2, 50)["total"], 120)
        with self.assertRaises(XlsxError):
            job.preview(0, 1000)
        with self.assertRaises(XlsxError):
            job.add({"slot": -1})
        with self.assertRaises(XlsxError):
            job.correct({"id": 0, "category": "not a prefix"})

    def test_mapping_wrong_modified_duplicate_and_unknown(self):
        job = self.job(); data = job.export(fake_batch)
        with self.assertRaises(XlsxError):
            restore(data + b"changed", job.dictionary)
        for edit in [lambda d: d["edits"].append(copy.deepcopy(d["edits"][0])),
                     lambda d: d["edits"][0].update(expected="unknown"),
                     lambda d: d.update(namespace="a" * 16),
                     lambda d: d.update(version=22)]:
            bad = copy.deepcopy(job.dictionary); edit(bad)
            with self.assertRaises(XlsxError):
                restore(data, bad)

    def test_residual_detection_blocks_unprocessed_known_value(self):
        # Simulated detector missing one occurrence on purpose. Known-value scan
        # must catch it, independently of the second detector pass.
        calls = 0
        def miss(texts):
            nonlocal calls
            results = fake_batch(texts)
            for result in results:
                if result:
                    calls += 1
                    if calls > 1:
                        result.clear()
            return results
        job = WorkbookJob(save(workbook({"A1": EMAIL, "A2": EMAIL})))
        job.analyze(miss, {})
        with self.assertRaisesRegex(XlsxError, "residui"):
            job.export(lambda xs: [[] for _ in xs])
        self.assertIsNone(job.output)

    def test_input_immutable_repeated_export(self):
        job = self.job()
        before = job.package.serialize()
        first = job.export(fake_batch)
        self.assertEqual(before, job.package.serialize())
        second = job.export(fake_batch)
        self.assertEqual(load_workbook(io.BytesIO(first)).active["A2"].value,
                         load_workbook(io.BytesIO(second)).active["A2"].value)


if __name__ == "__main__":
    unittest.main()
