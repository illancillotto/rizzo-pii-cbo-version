"""Conservative, in-memory OOXML processing; never evaluates formulas.

Only an explicit package subset is accepted. XML is edited directly to preserve
styles and layout; openpyxl is an independent structural reader after export.
All errors deliberately omit file content and user-supplied part names.
"""
from __future__ import annotations

import copy
import hashlib
import io
import json
import os
import posixpath
import re
import time
import zipfile
from collections import deque
from dataclasses import dataclass
from xml.etree import ElementTree as ET

from defusedxml.ElementTree import fromstring
from openpyxl import load_workbook
from openpyxl.formula.tokenizer import Tokenizer, TokenizerError
from openpyxl.utils.cell import range_boundaries
from openpyxl.writer.theme import theme_xml

from pii_session import PiiSession

MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL = "http://schemas.openxmlformats.org/package/2006/relationships"
OREL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
CT = "http://schemas.openxmlformats.org/package/2006/content-types"
NS = {"m": MAIN, "r": REL}
for prefix, uri in (("", MAIN), ("r", OREL)):
    ET.register_namespace(prefix, uri)


class XlsxError(ValueError):
    pass


@dataclass(frozen=True)
class Limits:
    compressed: int = 10 * 1024 * 1024
    expanded: int = 50 * 1024 * 1024
    parts: int = 500
    sheets: int = 32
    cells: int = 50000
    cell_length: int = 32767
    characters: int = 2000000
    seconds: int = 120

    @classmethod
    def environment(cls):
        values = {}
        for name in cls.__dataclass_fields__:
            raw = os.environ.get("PII_XLSX_" + name.upper())
            if raw is not None:
                try:
                    values[name] = int(raw)
                    if values[name] <= 0:
                        raise ValueError
                except ValueError:
                    raise XlsxError("Limite XLSX di configurazione non valido.") from None
        return cls(**values)


PART = re.compile(
    r"(?:\[Content_Types\]\.xml|_rels/\.rels|docProps/(?:core|app)\.xml|"
    r"xl/(?:workbook|styles|sharedStrings)\.xml|xl/theme/theme[0-9]+\.xml|"
    r"xl/worksheets/sheet[0-9]+\.xml|xl/(?:comments[0-9]+|comments/comment[0-9]+)\.xml|"
    r"xl/drawings/(?:commentsDrawing|vmlDrawing)[0-9]+\.vml|"
    r"xl/_rels/workbook\.xml\.rels|xl/worksheets/_rels/sheet[0-9]+\.xml\.rels)"
)
CONTENT_TYPES = {
    "application/vnd.openxmlformats-package.relationships+xml", "application/xml",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.styles+xml",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml",
    "application/vnd.openxmlformats-officedocument.theme+xml",
    "application/vnd.openxmlformats-package.core-properties+xml",
    "application/vnd.openxmlformats-officedocument.extended-properties+xml",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.comments+xml",
    "application/vnd.openxmlformats-officedocument.vmlDrawing",
}
REL_TYPES = {OREL + "/" + name for name in (
    "officeDocument", "worksheet", "styles", "theme", "sharedStrings",
    "extended-properties", "comments", "vmlDrawing",
)} | {"http://schemas.openxmlformats.org/package/2006/relationships/metadata/core-properties"}
BLOCKED_TAGS = {
    "extLst", "AlternateContent", "externalReferences", "externalReference",
    "pivotTable", "pivotCaches", "pivotTableParts", "pivotTablePart", "tableParts",
    "drawing", "picture", "oleObjects", "controls", "webPublishItems", "webPublishing",
    "dataValidations", "conditionalFormatting", "scenarios", "customSheetViews",
    "customWorkbookViews", "phoneticPr", "rPh", "smartTags", "cellSmartTags",
    "customPr", "customProperties", "sheetProtection", "workbookProtection",
}
SHEET_TAGS = set("worksheet sheetPr tabColor outlinePr pageSetUpPr dimension sheetViews sheetView "
                 "pane selection sheetFormatPr cols col sheetData row c f v is t r rPr "
                 "b i u strike outline shadow condense extend sz color rFont family charset "
                 "scheme vertAlign protection alignment sheetCalcPr autoFilter filterColumn "
                 "filters filter customFilters customFilter dynamicFilter top10 colorFilter "
                 "iconFilter sortState sortCondition mergeCells mergeCell printOptions "
                 "pageMargins pageSetup headerFooter oddHeader oddFooter evenHeader evenFooter "
                 "firstHeader firstFooter rowBreaks colBreaks brk legacyDrawing ignoredErrors "
                 "ignoredError sheetProtection".split())
CELL_REF = re.compile(r"\$?[A-Z]{1,3}\$?[1-9][0-9]*(?::\$?[A-Z]{1,3}\$?[1-9][0-9]*)?")
PLACEHOLDER = re.compile(r"\[\[RZ_[a-f0-9]{16}_[A-Z][A-Z0-9_]{0,31}_[1-9][0-9]*\]\]")
NEUTRAL = {"metadata_date": "2000-01-01T00:00:00Z", "metadata_number": "0", "metadata_bool": "false"}


def local(tag):
    return tag.rsplit("}", 1)[-1]


def text_of(node):
    return "".join(n.text or "" for n in node.iter() if local(n.tag) == "t")


def xml_bytes(node):
    return ET.tostring(node, encoding="utf-8", xml_declaration=True)


def contains_value(field, value):
    # Avoid identifying numeric 1 inside cell coordinate A1 or color 1F497D.
    # Exact single-character values are still checked in real data fields.
    return bool(re.search(r"(?<!\w)" + re.escape(value) + r"(?!\w)", PLACEHOLDER.sub(" ", field)))


class ValueScanner:
    """Aho-Corasick index: avoid values × fields quadratic residual scans."""
    def __init__(self, values, budget):
        self.next, self.fail, self.ends = [{}], [0], [[]]
        self.budget = budget
        for number, value in enumerate(values):
            if number % 256 == 0:
                budget()
            state = 0
            for char in value:
                if char not in self.next[state]:
                    self.next[state][char] = len(self.next)
                    self.next.append({}); self.fail.append(0); self.ends.append([])
                state = self.next[state][char]
            self.ends[state].append(value)
        queue = deque(self.next[0].values())
        while queue:
            state = queue.popleft()
            for char, child in self.next[state].items():
                queue.append(child)
                fallback = self.fail[state]
                while fallback and char not in self.next[fallback]:
                    fallback = self.fail[fallback]
                self.fail[child] = self.next[fallback].get(char, 0)
                self.ends[child].extend(self.ends[self.fail[child]])

    def find(self, field):
        field = PLACEHOLDER.sub(" ", field)
        state = 0
        for index, char in enumerate(field):
            if index % 4096 == 0:
                self.budget()
            while state and char not in self.next[state]:
                state = self.fail[state]
            state = self.next[state].get(char, 0)
            for value in self.ends[state]:
                start, end = index - len(value) + 1, index + 1
                if ((start == 0 or not (field[start - 1].isalnum() or field[start - 1] == "_")) and
                        (end == len(field) or not (field[end].isalnum() or field[end] == "_"))):
                    yield value


def parse_xml(data):
    try:
        return fromstring(data, forbid_dtd=True, forbid_entities=True, forbid_external=True)
    except Exception:
        raise XlsxError("XML non valido o non sicuro nel pacchetto XLSX.") from None


class Package:
    def __init__(self, data, filename="documento.xlsx", limits=None):
        self.limits = limits or Limits.environment()
        self.deadline = time.monotonic() + self.limits.seconds
        if not filename.lower().endswith(".xlsx"):
            raise XlsxError("Supportato solo .xlsx: .xls, .xlsb e .xlsm non sono accettati.")
        if data[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
            raise XlsxError("File cifrato o formato Excel legacy: salvare una copia .xlsx non cifrata.")
        if len(data) > self.limits.compressed:
            raise XlsxError("Superato il limite di dimensione compressa XLSX.")
        self.roots = {}
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                infos = archive.infolist()
                if (len(infos) > self.limits.parts or
                        sum(i.file_size for i in infos) > self.limits.expanded):
                    raise XlsxError("Superato il limite di parti o dimensione ZIP decompressa.")
                names = [i.filename for i in infos]
                if len(set(names)) != len(names):
                    raise XlsxError("Parti ZIP duplicate.")
                for info in infos:
                    self.check_time()
                    if info.flag_bits & 1:
                        raise XlsxError("File ZIP cifrato non supportato.")
                    if (not PART.fullmatch(info.filename) or info.compress_type not in
                            (zipfile.ZIP_STORED, zipfile.ZIP_DEFLATED)):
                        raise XlsxError("Componente non supportato: macro, immagini/oggetti, pivot, "
                                        "connessioni, tabelle o altra parte non ispezionabile.")
                    raw = archive.read(info)
                    if len(raw) != info.file_size:
                        raise XlsxError("Dimensione ZIP non valida.")
                    self.roots[info.filename] = parse_xml(raw)
        except XlsxError:
            raise
        except Exception:
            raise XlsxError("Il contenuto non è un pacchetto XLSX valido.") from None
        required = {"[Content_Types].xml", "_rels/.rels", "xl/workbook.xml"}
        if not required <= self.roots.keys():
            raise XlsxError("Parti obbligatorie XLSX mancanti.")
        self._preflight()
        self._expand_strings()
        self.sheets = []
        self.slots = []
        self.cells = {}
        self.formulas = []
        self.names = []
        self._discover()

    def check_time(self):
        if time.monotonic() > self.deadline:
            raise XlsxError("Budget di elaborazione XLSX esaurito.")

    def _preflight(self):
        types = self.roots["[Content_Types].xml"]
        if types.tag != "{" + CT + "}Types":
            raise XlsxError("Tipi di contenuto XLSX non validi.")
        for item in types:
            if item.get("ContentType") not in CONTENT_TYPES:
                raise XlsxError("Tipo di contenuto attivo o non supportato nel file XLSX.")
        if not any(n.get("PartName") == "/xl/workbook.xml" and n.get("ContentType") ==
                   "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml" for n in types):
            raise XlsxError("Il contenuto effettivo non è una cartella .xlsx senza macro.")
        for part, root in self.roots.items():
            if part.endswith(".rels"):
                if root.tag != "{" + REL + "}Relationships":
                    raise XlsxError("Relazioni XLSX non valide.")
                ids = set()
                for rel in root:
                    if rel.get("Id") in ids or not re.fullmatch(r"(?:rId[0-9]+|comments|anysvml)", rel.get("Id", "")):
                        raise XlsxError("Identificativi di relazione duplicati.")
                    ids.add(rel.get("Id"))
                    if rel.get("TargetMode") == "External" or rel.get("Type") not in REL_TYPES:
                        raise XlsxError("Collegamenti esterni, hyperlink e relazioni non supportate "
                                        "devono essere rimossi prima del caricamento.")
                    target = self.target(part, rel.get("Target", ""))
                    if target not in self.roots:
                        raise XlsxError("Relazione verso una parte mancante.")
            for node in root.iter():
                if (node.tail or "").strip():
                    raise XlsxError("Testo XML fuori dai campi supportati.")
                if local(node.tag) in BLOCKED_TAGS and not (
                        local(node.tag) == "workbookProtection" and not node.attrib):
                    raise XlsxError("Funzionalità XLSX non supportata: estensioni, protezione, "
                                    "convalide, filtri complessi o componenti attivi.")
                if len(node.text or "") > self.limits.cell_length:
                    raise XlsxError("Superata la lunghezza massima di un campo XLSX.")
                if any(len(v) > self.limits.cell_length for v in node.attrib.values()):
                    raise XlsxError("Attributo XLSX troppo lungo.")
            if re.fullmatch(r"xl/worksheets/sheet[0-9]+\.xml", part):
                if root.tag != "{" + MAIN + "}worksheet":
                    raise XlsxError("Schema del foglio non supportato.")
                if any(local(n.tag) not in SHEET_TAGS or
                       not n.tag.startswith("{" + MAIN + "}") for n in root.iter()):
                    raise XlsxError("Elemento del foglio non supportato (inclusi hyperlink interni/esterni).")
            if part.endswith(".vml"):
                # Legacy comments only; no images, arbitrary VML text or controls.
                allowed = set("xml shapelayout idmap shapetype stroke path shape fill shadow "
                              "textbox div ClientData MoveWithCells SizeWithCells Anchor AutoFill "
                              "Row Column Visible lock".split())
                if any(local(n.tag) not in allowed for n in root.iter()):
                    raise XlsxError("Disegno VML non limitato a commenti semplici.")
                clients = [n for n in root.iter() if local(n.tag) == "ClientData"]
                if not clients or any(n.get("ObjectType") != "Note" for n in clients):
                    raise XlsxError("Controlli VML non supportati.")
                if any((n.text or "").strip() for n in root.iter()
                       if local(n.tag) not in {"Anchor", "Row", "Column", "AutoFill"}):
                    raise XlsxError("Testo non ispezionabile in disegni VML.")
            if part.startswith("xl/theme/"):
                # The first version accepts the audited Office theme shipped by
                # openpyxl. Arbitrary themes can contain text/drawings/extensions.
                if xml_bytes(root) != xml_bytes(parse_xml(theme_xml.encode())):
                    raise XlsxError("Tema Excel personalizzato non supportato: usare il tema Office standard.")

    @staticmethod
    def target(rels_part, target):
        if target.startswith("/"):
            return target[1:]
        folder = posixpath.dirname(rels_part).replace("/_rels", "")
        if folder == "_rels":
            folder = ""
        return posixpath.normpath(posixpath.join(folder, target))

    def _expand_strings(self):
        shared = self.roots.get("xl/sharedStrings.xml")
        strings = list(shared) if shared is not None else []
        for part, root in self.roots.items():
            if not re.fullmatch(r"xl/worksheets/sheet[0-9]+\.xml", part):
                continue
            for cell in root.findall(".//m:c", NS):
                if cell.get("t") == "s":
                    try:
                        index = int(cell.find("m:v", NS).text)
                        if index < 0:
                            raise ValueError
                        item = copy.deepcopy(strings[index])
                    except Exception:
                        raise XlsxError("Indice di stringa condivisa non valido.") from None
                    item.tag = "{" + MAIN + "}is"
                    cell.remove(cell.find("m:v", NS))
                    cell.set("t", "inlineStr")
                    cell.append(item)
        if shared is not None:
            del self.roots["xl/sharedStrings.xml"]
            for node in list(self.roots["[Content_Types].xml"]):
                if node.get("PartName") == "/xl/sharedStrings.xml":
                    self.roots["[Content_Types].xml"].remove(node)
            for part, root in self.roots.items():
                if part.endswith(".rels"):
                    for node in list(root):
                        if node.get("Type") == OREL + "/sharedStrings":
                            root.remove(node)

    def _slot(self, part, node, value, kind="text", **meta):
        if len(value) > self.limits.cell_length:
            raise XlsxError("Superata la lunghezza massima di una cella.")
        if re.search(r"_x[0-9a-fA-F]{4}_", value):
            raise XlsxError("Sequenze di escape OOXML nei dati non supportate: salvare una copia con testo semplice.")
        slot = {"id": len(self.slots), "part": part, "node": node, "value": value,
                "kind": kind, **meta}
        self.slots.append(slot)
        return slot

    def _discover(self):
        workbook = self.roots["xl/workbook.xml"]
        if workbook.tag != "{" + MAIN + "}workbook":
            raise XlsxError("Schema della cartella non supportato.")
        rels = self.roots.get("xl/_rels/workbook.xml.rels")
        if rels is None:
            raise XlsxError("Relazioni della cartella mancanti.")
        targets = {r.get("Id"): self.target("xl/_rels/workbook.xml.rels", r.get("Target", ""))
                   for r in rels if r.get("Type") == OREL + "/worksheet"}
        sheets = workbook.findall("m:sheets/m:sheet", NS)
        if not 0 < len(sheets) <= self.limits.sheets:
            raise XlsxError("Superato il limite di fogli o cartella vuota.")
        processed, count, area = set(), 0, 0
        for index, sheet in enumerate(sheets):
            part = targets.get(sheet.get("{" + OREL + "}id"))
            if part not in self.roots or part in processed:
                raise XlsxError("Riferimenti dei fogli non validi.")
            processed.add(part)
            title = sheet.get("name", "")
            self.names.append(title)
            root = self.roots[part]
            dim = root.find("m:dimension", NS)
            try:
                bounds = range_boundaries(dim.get("ref")) if dim is not None else (1, 1, 1, 1)
                area += bounds[2] * bounds[3]
            except Exception:
                raise XlsxError("Dimensioni del foglio non valide.") from None
            if area > self.limits.cells:
                raise XlsxError("Superato il limite dell'area complessiva dei fogli.")
            max_row = max_col = 0
            for cell in root.findall(".//m:c", NS):
                count += 1
                if count > self.limits.cells:
                    raise XlsxError("Superato il limite di celle.")
                coordinate = cell.get("r", "")
                if not re.fullmatch(r"[A-Z]{1,3}[1-9][0-9]*", coordinate):
                    raise XlsxError("Coordinata di cella non valida.")
                col, row, _, _ = range_boundaries(coordinate)
                max_row, max_col = max(max_row, row), max(max_col, col)
                if row * col > self.limits.cells or (index, coordinate) in self.cells:
                    raise XlsxError("Cella duplicata o fuori dal budget di area.")
                self.cells[index, coordinate] = cell
                formula = cell.find("m:f", NS)
                if formula is not None:
                    if formula.attrib:
                        raise XlsxError("Formule condivise, matriciali o con attributi non supportate.")
                    refs = self._formula(formula.text or "")
                    self.formulas.append((index, refs))
                    for cached in list(cell.findall("m:v", NS)):
                        cell.remove(cached)
                    continue
                kind = cell.get("t", "n")
                if kind == "inlineStr":
                    value = text_of(cell)
                    kind = "text"
                else:
                    value = cell.findtext("m:v", default="", namespaces=NS)
                    if kind == "str":
                        kind = "text"
                    if kind not in {"text", "n", "b", "d", "e"}:
                        raise XlsxError("Tipo di cella non supportato.")
                if value:
                    self._slot(part, cell, value, kind, sheet=index, cell=coordinate,
                               row=row, column=col, cell_type=cell.get("t"))
            self.sheets.append({"index": index, "name": title, "state": sheet.get("state", "visible"),
                                "rows": max_row, "columns": max_col})
            for node in root.iter():
                if local(node.tag) in {"oddHeader", "oddFooter", "evenHeader", "evenFooter",
                                       "firstHeader", "firstFooter"} and node.text:
                    self._slot(part, node, node.text, field=local(node.tag), sheet=index)
                # Filter values can contain old personal values. Inspect each one.
                if local(node.tag) in {"filter", "customFilter"} and node.get("val"):
                    self._slot(part, node, node.get("val"), field="filter", attribute="val", sheet=index)
        # Refuse orphan worksheets (including data not listed by the workbook).
        if {p for p in self.roots if re.fullmatch(r"xl/worksheets/sheet[0-9]+\.xml", p)} != processed:
            raise XlsxError("Fogli orfani non supportati.")
        for node in workbook.findall("m:definedNames/m:definedName", NS):
            self.names.append(node.get("name", ""))
            # Only static ranges. No literal values, expressions, DDE or external links.
            value = node.text or ""
            match = re.fullmatch(r"(?:'((?:[^']|'')+)'|([^'!]+))!(\$?[A-Z]+\$?[1-9][0-9]*(?::\$?[A-Z]+\$?[1-9][0-9]*)?)", value)
            if not match or (match.group(1) or match.group(2)).replace("''", "'") not in self.names:
                raise XlsxError("Nomi definiti: supportati solo riferimenti statici a intervalli locali.")
        for part, root in self.roots.items():
            if part.startswith("docProps/") or re.fullmatch(r"xl/(?:comments[0-9]+|comments/comment[0-9]+)\.xml", part):
                for node in root.iter():
                    if node.text and node.text.strip():
                        tag = local(node.tag)
                        kind = "text"
                        if part.startswith("docProps/"):
                            if tag in {"created", "modified", "lastPrinted", "filetime"}:
                                kind = "metadata_date"
                            elif tag in {"i2", "i4", "i8", "int", "ui1", "ui2", "ui4", "ui8", "uint", "r4", "r8", "decimal"}:
                                kind = "metadata_number"
                            elif tag == "bool":
                                kind = "metadata_bool"
                        self._slot(part, node, node.text, kind=kind, field=tag)
            if part == "xl/styles.xml":
                # Textual style metadata must be inspected too (custom number formats,
                # font names and named styles). Changing their semantics is unsafe.
                for node in root.iter():
                    for attr in ("formatCode", "name"):
                        if node.get(attr):
                            self.names.append(node.get(attr))
        if sum(len(s["value"]) for s in self.slots) > self.limits.characters:
            raise XlsxError("Superato il budget complessivo di caratteri XLSX.")
        calc = workbook.find("m:calcPr", NS)
        if calc is None:
            calc = ET.SubElement(workbook, "{" + MAIN + "}calcPr")
        calc.attrib.update(calcMode="auto", fullCalcOnLoad="1", forceFullCalc="1")
        self.static_fields = self._static_fields()
        if sum(map(len, self.static_fields)) + sum(len(s["value"]) for s in self.slots) > self.limits.characters:
            raise XlsxError("Superato il budget complessivo di caratteri XLSX.")
        self.check_time()

    def _static_fields(self):
        """All XML scalars not represented by editable data slots.

        This includes orphan comments, attributes, theme/style metadata and
        workbook metadata. They are checked and block export if sensitive; they
        are never silently ignored or blindly rewritten.
        """
        fields = []
        covered_text, covered_attrs = set(), set()
        for slot in self.slots:
            node = slot["node"]
            if slot.get("attribute"):
                covered_attrs.add((id(node), slot["attribute"]))
            elif "cell" in slot:
                covered_text.update(id(n) for n in node.iter())
            else:
                covered_text.add(id(node))
        for part, root in self.roots.items():
            if part.startswith("xl/theme/"):
                continue  # exact, preflight-verified technical boilerplate
            for node in root.iter():
                if id(node) not in covered_text and (node.text or "").strip():
                    fields.append(node.text)
                # Ignore only package/schema identifiers, whose values were
                # validated by preflight, and standard namespace URI constants.
                for attr, value in node.attrib.items():
                    if (id(node), attr) in covered_attrs:
                        continue
                    if part.endswith(".rels") or part == "[Content_Types].xml":
                        continue
                    technical = (re.fullmatch(r"[+-]?[0-9]+(?:\.[0-9]+)?", value) or
                                 CELL_REF.fullmatch(value) or re.fullmatch(r"rId[0-9]+", value) or
                                 (attr in {"rgb", "val"} and re.fullmatch(r"[A-Fa-f0-9]{6,8}", value)))
                    safe_font = local(node.tag) == "name" and attr == "val" and value in {
                        "Calibri", "Arial", "Times New Roman", "Cambria", "Helvetica", "Verdana",
                        "Courier New", "Segoe UI", "Tahoma", "Symbol"}
                    if value and not technical and not safe_font:
                        fields.append(value)
        return list(dict.fromkeys(fields))

    @staticmethod
    def _formula(formula):
        if len(formula) > 2048:
            raise XlsxError("Formula troppo lunga.")
        try:
            tokens = Tokenizer("=" + formula).items
        except TokenizerError:
            raise XlsxError("Formula non supportata.") from None
        refs = []
        for token in tokens:
            if token.type == "OPERAND":
                if token.subtype == "RANGE" and CELL_REF.fullmatch(token.value):
                    refs.append(range_boundaries(token.value.replace("$", "")))
                elif token.subtype != "NUMBER":
                    raise XlsxError("Formule con stringhe, nomi, riferimenti esterni o tra fogli non supportate.")
            elif token.type == "FUNC":
                if token.subtype == "OPEN" and token.value.upper() not in {
                        "SUM(", "MIN(", "MAX(", "AVERAGE(", "COUNT(", "ABS(", "ROUND("}:
                    raise XlsxError("Funzione Excel non supportata.")
            elif token.type not in {"OP_PRE", "OP_IN", "OP_POST", "PAREN", "SEP", "WHITE-SPACE"}:
                raise XlsxError("Formula non supportata.")
        if not tokens:
            raise XlsxError("Formula vuota non supportata.")
        return refs

    def serialize(self):
        stream = io.BytesIO()
        with zipfile.ZipFile(stream, "w", zipfile.ZIP_DEFLATED) as archive:
            for name, root in self.roots.items():
                archive.writestr(name, xml_bytes(root))
        return stream.getvalue()


class WorkbookJob:
    def __init__(self, data, filename="documento.xlsx", limits=None):
        self.package = Package(data, filename, limits)
        self.changes = []
        self.session = None
        self.excluded = set()
        self.output = None
        self.dictionary = None
        self.analysed = False
        self.progress = {"phase": "uploaded", "completed": 0, "total": len(self.package.slots)}

    def summary(self):
        return {"sheets": self.package.sheets, "cells": len(self.package.cells),
                "fields": len(self.package.slots), "formulas": len(self.package.formulas),
                "unsupported": [], "policy": "strict-xlsx-v1"}

    def analyze(self, detect_batch, config):
        p = self.package
        p.deadline = time.monotonic() + p.limits.seconds
        self.output = self.dictionary = None
        self.analysed = False
        self.changes = []
        self.progress = {"phase": "analysis", "completed": 0, "total": len(p.slots)}
        if not isinstance(config, dict):
            raise XlsxError("Configurazione XLSX non valida.")
        mapping = config.get("include_mapping", True)
        if not isinstance(mapping, bool):
            raise XlsxError("include_mapping deve essere booleano.")
        excluded = config.get("exclude_tags", [])
        if not isinstance(excluded, list) or any(not isinstance(t, str) for t in excluded):
            raise XlsxError("Categorie escluse non valide.")
        self.excluded = set(excluded)
        self.session = PiiSession([s["value"] for s in p.slots] + p.names + p.static_fields, mapping)
        selected = config.get("sheets", list(range(len(p.sheets))))
        if (not isinstance(selected, list) or any(type(i) is not int or i < 0 or
                                                i >= len(p.sheets) for i in selected)):
            raise XlsxError("Selezione fogli non valida.")
        selected = set(selected)
        rules = config.get("columns", [])
        if not isinstance(rules, list) or len(rules) > p.limits.cells:
            raise XlsxError("Regole di colonna non valide.")
        indexed_rules = {}
        for rule in rules:
            if not isinstance(rule, dict):
                raise XlsxError("Regola di colonna non valida.")
            sheet, column = rule.get("sheet"), rule.get("column")
            header, start = rule.get("header_row", 1), rule.get("start_row", 2)
            category = rule.get("category", "ID")
            if (type(sheet) is not int or sheet not in selected or type(column) is not int or
                    not 1 <= column <= p.sheets[sheet]["columns"] or
                    type(header) is not int or type(start) is not int or header < 0 or start < 1 or
                    not isinstance(category, str) or not re.fullmatch(r"[A-Z][A-Z0-9_]{0,31}", category)):
                raise XlsxError("Selezione foglio/colonna, righe o categoria non valida.")
            if (sheet, column) in indexed_rules:
                raise XlsxError("Regole duplicate per una colonna.")
            indexed_rules[sheet, column] = (header, start, category)
        # Sheet and defined names and style metadata cannot be changed safely:
        # detect with no exclusions, then fail closed if personal data is found.
        checked_names = list(dict.fromkeys(p.names + p.static_fields))
        for offset in range(0, len(checked_names), 16):
            p.check_time()
            if any(detect_batch(checked_names[offset:offset + 16])):
                raise XlsxError("Dati personali in nomi di fogli/nomi definiti/stili: rinominarli "
                                "prima di elaborare il documento.")
        automatic = config.get("automatic", True)
        if not isinstance(automatic, bool):
            raise XlsxError("automatic deve essere booleano.")
        slots = p.slots
        for offset in range(0, len(slots), 16):
            p.check_time()
            batch = slots[offset:offset + 16]
            # Each cell is a separate sample: no entity can cross a cell boundary.
            detections = detect_batch([s["value"] for s in batch])
            if len(detections) != len(batch):
                raise XlsxError("Risposta del motore di riconoscimento non valida.")
            for slot, entities in zip(batch, detections):
                if "cell" in slot and slot["sheet"] not in selected:
                    # Explicitly unselected cells are still checked; never export
                    # silently detected PII in an unprocessed hidden sheet.
                    if any(e["label"] not in self.excluded for e in entities):
                        raise XlsxError("Dati personali rilevati in un foglio non selezionato.")
                    continue
                rule = indexed_rules.get((slot.get("sheet"), slot.get("column")))
                if rule and slot["row"] >= rule[1] and slot["row"] != rule[0]:
                    entities = [{"start": 0, "end": len(slot["value"]), "label": rule[2], "forced": True}]
                elif not automatic and "cell" in slot:
                    entities = []
                else:
                    entities = [e for e in entities if e["label"] not in self.excluded]
                for entity in entities:
                    start, end, category = entity["start"], entity["end"], entity["label"]
                    if not 0 <= start < end <= len(slot["value"]):
                        raise XlsxError("Offset del rilevatore non valido.")
                    if slot["kind"] != "text":
                        # Numeric identifiers become text pseudonyms; their exact
                        # lexical value and Excel type are restored by the dictionary.
                        start, end = 0, len(slot["value"])
                    if any(c["slot"] == slot["id"] and c["start"] < end and start < c["end"]
                           for c in self.changes):
                        continue
                    ph = self.session.placeholder(category, slot["value"][start:end], slot["kind"])
                    ph = NEUTRAL.get(slot["kind"], ph)
                    self.changes.append({"id": len(self.changes), "slot": slot["id"],
                                         "start": start, "end": end, "category": category,
                                         "replacement": ph, "enabled": True,
                                         "forced": bool(entity.get("forced")),
                                         "conflict": bool(entity.get("forced") and category in self.excluded)})
            self.progress = {"phase": "analysis", "completed": min(offset + 16, len(slots)), "total": len(slots)}
        p.check_time()
        self.analysed = True
        self.progress["phase"] = "review"
        return {"changes": len(self.changes), "mapping_enabled": mapping,
                "forced_precedence": True, "complete": True}

    def preview(self, page=0, size=50):
        if not self.analysed:
            raise XlsxError("Analizzare prima il documento.")
        if page < 0 or not 1 <= size <= 100:
            raise XlsxError("Pagina non valida.")
        result = []
        for change in self.changes[page * size:(page + 1) * size]:
            slot = self.package.slots[change["slot"]]
            item = {k: v for k, v in change.items() if k != "slot"}
            item.update({"slot": slot["id"], "sheet": slot.get("sheet"),
                         "cell": slot.get("cell"), "field": slot.get("field", "cell")})
            if self.session.mapping_enabled:
                item["original"] = slot["value"][change["start"]:change["end"]]
            result.append(item)
        return {"page": page, "size": size, "total": len(self.changes), "items": result}

    def correct(self, edit):
        if not self.analysed or not isinstance(edit, dict):
            raise XlsxError("Revisione non valida.")
        identifier = edit.get("id")
        if type(identifier) is not int or not 0 <= identifier < len(self.changes):
            raise XlsxError("Sostituzione non valida.")
        change = self.changes[identifier]
        if edit.get("scope", "occurrence") not in {"occurrence", "associated"}:
            raise XlsxError("Ambito di correzione non valido.")
        targets = [change] if edit.get("scope", "occurrence") == "occurrence" else [
            c for c in self.changes if c["replacement"] == change["replacement"]]
        if "enabled" in edit and not isinstance(edit["enabled"], bool):
            raise XlsxError("enabled deve essere booleano.")
        if "enabled" in edit and not edit["enabled"] and any(c["forced"] for c in targets):
            raise XlsxError("Una colonna obbligatoria non può essere esclusa in revisione.")
        if "category" in edit:
            slot = self.package.slots[change["slot"]]
            try:
                ph = self.session.placeholder(edit["category"], slot["value"][change["start"]:change["end"]], slot["kind"])
                ph = NEUTRAL.get(slot["kind"], ph)
            except ValueError:
                raise XlsxError("Categoria non valida.") from None
            for target in targets:
                target.update(category=edit["category"], replacement=ph)
        for target in targets:
            if "enabled" in edit:
                target["enabled"] = edit["enabled"]
        self.output = self.dictionary = None

    def add(self, edit):
        if not self.analysed or not isinstance(edit, dict):
            raise XlsxError("Aggiunta non valida.")
        slot_id = edit.get("slot")
        if slot_id is None and type(edit.get("sheet")) is int and isinstance(edit.get("cell"), str):
            slot_id = next((s["id"] for s in self.package.slots if s.get("sheet") == edit["sheet"]
                            and s.get("cell") == edit["cell"]), None)
        if type(slot_id) is not int or not 0 <= slot_id < len(self.package.slots):
            raise XlsxError("Campo o cella non valido.")
        slot = self.package.slots[slot_id]
        start, end = edit.get("start", 0), edit.get("end", len(slot["value"]))
        if (type(start) is not int or type(end) is not int or not 0 <= start < end <= len(slot["value"]) or
                (slot["kind"] != "text" and (start != 0 or end != len(slot["value"])))):
            raise XlsxError("Intervallo non valido.")
        if any(c["slot"] == slot_id and c["enabled"] and c["start"] < end and start < c["end"]
               for c in self.changes):
            raise XlsxError("Sostituzioni sovrapposte: escludere prima la proposta esistente.")
        category = edit.get("category", "ID")
        try:
            ph = self.session.placeholder(category, slot["value"][start:end], slot["kind"])
            ph = NEUTRAL.get(slot["kind"], ph)
        except ValueError:
            raise XlsxError("Categoria non valida.") from None
        self.changes.append({"id": len(self.changes), "slot": slot_id, "start": start, "end": end,
                             "category": category, "replacement": ph, "enabled": True,
                             "forced": False, "conflict": False})
        self.output = self.dictionary = None

    @staticmethod
    def set_value(slot, value, kind="text", cell_type=None):
        node = slot["node"]
        if "cell" not in slot:
            if slot.get("attribute"):
                node.set(slot["attribute"], value)
            else:
                node.text = value
            return
        for child in list(node):
            if local(child.tag) in {"v", "is"}:
                node.remove(child)
        if kind == "text":
            node.set("t", "inlineStr")
            inline = ET.SubElement(node, "{" + MAIN + "}is")
            text = ET.SubElement(inline, "{" + MAIN + "}t")
            text.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
            text.text = value
        else:
            if cell_type is None:
                node.attrib.pop("t", None)
            else:
                node.set("t", cell_type)
            ET.SubElement(node, "{" + MAIN + "}v").text = value

    def export(self, detect_batch):
        if not self.analysed:
            raise XlsxError("Analizzare prima il documento.")
        p = self.package
        self.output = self.dictionary = None
        self.progress = {"phase": "verification", "completed": 0, "total": len(p.slots)}
        p.deadline = time.monotonic() + p.limits.seconds
        # Always export from the analysed source, so repeated exports and review
        # changes cannot compound substitutions.
        working = copy.deepcopy(p)
        active = [c for c in self.changes if c["enabled"]]
        modified_cells = { (p.slots[c["slot"]].get("sheet"), p.slots[c["slot"]].get("cell"))
                           for c in active if "cell" in p.slots[c["slot"]] }
        for sheet, refs in p.formulas:
            for row_sheet, coordinate in modified_cells:
                if row_sheet != sheet:
                    continue
                col, row, _, _ = range_boundaries(coordinate)
                if any(a <= col <= c and b <= row <= d for a, b, c, d in refs):
                    raise XlsxError("Una formula dipende da celle da mascherare. Esportare prima "
                                    "una copia con formule convertite in valori.")
        by_slot = {}
        for change in active:
            by_slot.setdefault(change["slot"], []).append(change)
        restoration = []
        known_values = set()
        for slot_id, changes in by_slot.items():
            slot = working.slots[slot_id]
            old = slot["value"]
            value = old
            for change in sorted(changes, key=lambda c: c["start"], reverse=True):
                known_values.add((old[change["start"]:change["end"]], slot["kind"]))
                value = value[:change["start"]] + change["replacement"] + value[change["end"]:]
            if len(value) > p.limits.cell_length:
                raise XlsxError("La sostituzione supera la lunghezza massima della cella.")
            if self.session.mapping_enabled:
                restoration.append({"slot": slot_id, "expected": value, "original": old,
                                    "kind": slot["kind"], "cell_type": slot.get("cell_type")})
            self.set_value(slot, value)
        data = working.serialize()
        p.check_time()
        verified = Package(data, limits=p.limits)
        self._verify_structure(data, p, verified)
        scanner = ValueScanner({v for v, k in known_values if k == "text"}, p.check_time)
        typed_values = {v for v, k in known_values if k != "text"}
        for field in verified.static_fields:
            if field in typed_values or next(scanner.find(field), None) is not None:
                raise XlsxError("Verifica residui fallita nei nomi, stili o metadati non modificabili.")
        # Scan every accepted XML scalar, including metadata, relations and caches.
        # Review exclusions authorize only the exact reviewed field/occurrence.
        authorised = {}
        for change in self.changes:
            if not change["enabled"]:
                slot = p.slots[change["slot"]]
                authorised.setdefault(slot["id"], set()).add(slot["value"][change["start"]:change["end"]])
        for slot in verified.slots:
            value = slot["value"]
            source_slot = p.slots[slot["id"]] if slot["id"] < len(p.slots) else None
            allowed = authorised.get(slot["id"], set())
            for original in scanner.find(value):
                if original not in allowed:
                    raise XlsxError("Verifica residui fallita: un valore sostituito è ancora presente.")
            if value in typed_values and value not in allowed:
                raise XlsxError("Verifica residui fallita: un identificativo tipizzato è ancora presente.")
            if source_slot is None:
                raise XlsxError("Campi del documento non coerenti dopo l'esportazione.")
        # Run the same detector again on output text, removing our placeholders.
        for offset in range(0, len(verified.slots), 16):
            p.check_time()
            batch = verified.slots[offset:offset + 16]
            clean = ["" if s["kind"] in NEUTRAL and s["value"] == NEUTRAL[s["kind"]]
                     else PLACEHOLDER.sub(" ", s["value"]) for s in batch]
            results = detect_batch(clean)
            for slot, value, entities in zip(batch, clean, results):
                for entity in entities:
                    original = value[entity["start"]:entity["end"]]
                    if entity["label"] not in self.excluded and original not in authorised.get(slot["id"], set()):
                        raise XlsxError("Verifica residui fallita: rivedere le sostituzioni prima del download.")
            self.progress = {"phase": "verification", "completed": min(offset + 16, len(verified.slots)),
                             "total": len(verified.slots)}
        self.output = data
        if self.session.mapping_enabled:
            self.dictionary = {"format": "rizzo-pii-xlsx", "version": 1,
                               "sha256": hashlib.sha256(data).hexdigest(),
                               "namespace": self.session.namespace, "edits": restoration}
        self.progress["phase"] = "ready"
        return data

    @staticmethod
    def _verify_structure(data, source, output):
        if source.sheets != output.sheets or len(source.cells) != len(output.cells):
            raise XlsxError("Verifica strutturale XLSX fallita.")
        try:
            book = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
            if len(book.worksheets) != len(source.sheets):
                raise ValueError
            book.close()
        except Exception:
            raise XlsxError("Il file esportato non supera la riapertura indipendente.") from None


def restore(data, dictionary, limits=None):
    limits = limits or Limits.environment()
    if not isinstance(dictionary, dict) or set(dictionary) != {"format", "version", "sha256", "namespace", "edits"}:
        raise XlsxError("Formato del dizionario XLSX non valido.")
    if (dictionary["format"] != "rizzo-pii-xlsx" or type(dictionary["version"]) is not int or dictionary["version"] != 1 or
            not isinstance(dictionary["namespace"], str) or not re.fullmatch(r"[a-f0-9]{16}", dictionary["namespace"]) or
            dictionary["sha256"] != hashlib.sha256(data).hexdigest()):
        raise XlsxError("Dizionario errato o file modificato/rielaborato: SHA-256 non corrispondente.")
    package = Package(data, limits=limits)
    edits = dictionary["edits"]
    if not isinstance(edits, list) or len(edits) > len(package.slots):
        raise XlsxError("Elenco di ripristino non valido.")
    seen = set()
    for edit in edits:
        if not isinstance(edit, dict) or set(edit) != {"slot", "expected", "original", "kind", "cell_type"}:
            raise XlsxError("Voce di ripristino non valida.")
        slot_id = edit["slot"]
        if type(slot_id) is not int or not 0 <= slot_id < len(package.slots) or slot_id in seen:
            raise XlsxError("Campo di ripristino duplicato o sconosciuto.")
        seen.add(slot_id)
        slot = package.slots[slot_id]
        if (not isinstance(edit["original"], str) or len(edit["original"]) > limits.cell_length or
                not isinstance(edit["kind"], str) or
                (edit["cell_type"] is not None and not isinstance(edit["cell_type"], str)) or
                edit["expected"] != slot["value"] or edit["kind"] not in {"text", "n", "b", "d", "e", *NEUTRAL} or
                edit["cell_type"] not in {None, "n", "b", "d", "e", "str", "inlineStr"}):
            raise XlsxError("Valore o tipo di ripristino non valido.")
        matches = PLACEHOLDER.findall(edit["expected"])
        neutral = edit["kind"] in NEUTRAL and slot["kind"] == edit["kind"] and edit["expected"] == NEUTRAL[edit["kind"]]
        added = set(matches) - set(PLACEHOLDER.findall(edit["original"]))
        if not neutral and (not added or any("RZ_" + dictionary["namespace"] + "_" not in ph for ph in added)):
            raise XlsxError("Placeholder sconosciuto o non coerente con il dizionario.")
        WorkbookJob.set_value(slot, edit["original"], edit["kind"], edit["cell_type"])
    own_marker = "[[RZ_" + dictionary["namespace"] + "_"
    for slot in package.slots:
        if own_marker in slot["value"] and slot["id"] not in seen:
            raise XlsxError("Dizionario incompleto: placeholder privi di una voce di ripristino.")
    result = package.serialize()
    checked = Package(result, limits=limits)
    WorkbookJob._verify_structure(result, package, checked)
    return result
