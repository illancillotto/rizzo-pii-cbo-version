#!/usr/bin/env python3
"""Offline synthetic XLSX end-to-end check with actual local model weights.

Not a benchmark. Never patches the detector or downloads a model. Run after
installing app requirements and setting PII_MODEL_DIR to an existing model.
"""
import io
import json
import os
import sys
import time
import zipfile
from pathlib import Path

os.environ["HF_HUB_OFFLINE"] = "1"
os.environ["TRANSFORMERS_OFFLINE"] = "1"
os.environ.setdefault("OMP_NUM_THREADS", "4")
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src/app"))

import app  # noqa: E402 — real pipeline, no mocks
from openpyxl import Workbook, load_workbook  # noqa: E402
from openpyxl.comments import Comment  # noqa: E402

HEADER = {"X-Rizzo-XLSX": "1"}
EMAIL = "sintetico@example.invalid"
NAME = "Mario Rossi"  # synthetic fixture identity, not a real person
CF = "RSSMRA85H12F205Y"
PIVA = "12345678903"
IBAN = "IT60X0542811101000000123456"


def require(condition, reason):
    if not condition:
        raise AssertionError(reason)  # reasons contain no document content


def wait(client, url):
    deadline = time.monotonic() + 150
    while time.monotonic() < deadline:
        response = client.get(url + "/status")
        require(response.status_code == 200, "Status API failed")
        state = response.json
        if state["state"] not in {"analysing", "exporting"}:
            require(not state["error"], state["error"] or "Background operation failed")
            return state
        time.sleep(.05)
    raise AssertionError("Local inference timed out")


def main():
    sample = f"Il cliente {NAME}, C.F. {CF}, P.IVA {PIVA}, IBAN {IBAN}, email {EMAIL}."
    labels = {e["label"] for e in app.recognize_batch([sample])[0]}
    require({"FULLNAME", "CF", "PIVA", "IBAN", "EMAIL"} <= labels,
            "Real model + regex did not detect all five fixture categories")
    book = Workbook()
    sheet = book.active
    for cell, value in {"A1": "Testo", "A2": sample, "B1": "Matricola", "B2": 7654321,
                        "C1": "Risultato", "C2": "=SUM(D2:D3)", "D1": "Quantita", "D2": 4, "D3": 5}.items():
        sheet[cell] = value
    book.create_sheet("Secondo")["A1"] = EMAIL
    book.worksheets[1].sheet_state = "veryHidden"
    sheet.row_dimensions[2].hidden = True
    sheet["A2"].comment = Comment("Nota " + EMAIL, NAME)
    sheet.oddHeader.center.text = EMAIL
    book.properties.creator = NAME
    stream = io.BytesIO(); book.save(stream)
    client = app.app.test_client()
    upload = client.post("/api/xlsx/upload", headers=HEADER,
                         data={"file": (io.BytesIO(stream.getvalue()), "synthetic.xlsx")})
    require(upload.status_code == 201, "Upload failed")
    url = "/api/xlsx/" + upload.json["id"]
    try:
        response = client.post(url + "/analyze", headers=HEADER, json={
            "include_mapping": True, "exclude_tags": ["AMOUNT"],
            "columns": [{"sheet": 0, "column": 2, "header_row": 1, "start_row": 2, "category": "ACCOUNT"}]})
        require(response.status_code == 202, "Analysis failed to start")
        wait(client, url)
        preview = client.get(url + "/preview?size=100").json
        require(preview["total"] <= 100, "Fixture unexpectedly exceeded one review page")
        for item in preview["items"]:
            if item["cell"] in {"D2", "D3"}:
                response = client.patch(url + "/review", headers=HEADER,
                                        json={"id": item["id"], "enabled": False, "scope": "occurrence"})
                require(response.status_code == 200, "Measure review failed")
        contact = next(item for item in preview["items"] if item.get("original") == EMAIL)
        response = client.patch(url + "/review", headers=HEADER,
                                json={"id": contact["id"], "category": "CONTACT", "scope": "associated"})
        require(response.status_code == 200, "Associated review failed")
        require(client.post(url + "/export", headers=HEADER).status_code == 202, "Export failed to start")
        wait(client, url)
        output = client.get(url + "/download")
        mapping = client.get(url + "/mapping")
        require(output.status_code == mapping.status_code == 200, "Verified download failed")
        with zipfile.ZipFile(io.BytesIO(output.data)) as archive:
            for part in archive.namelist():
                raw = archive.read(part)
                require(not any(value.encode() in raw for value in [NAME, EMAIL, CF, PIVA, IBAN]),
                        "Known personal value remains in the exported package")
        masked = load_workbook(io.BytesIO(output.data))
        require(masked.worksheets[1]["A1"].value in masked.active["A2"].value, "Cross-sheet identity failed")
        require(masked.active["C2"].value == "=SUM(D2:D3)", "Formula preservation failed")
        require(load_workbook(io.BytesIO(output.data), data_only=True).active["C2"].value is None,
                "Formula cache was not cleared")
        response = client.post("/api/xlsx/restore", headers=HEADER,
            data={"file": (io.BytesIO(output.data), "masked.xlsx"),
                  "mapping": (io.BytesIO(mapping.data), "dictionary.json")})
        require(response.status_code == 200, "Restore failed")
        restored = load_workbook(io.BytesIO(response.data))
        checked = 0
        for source, destination in zip(book.worksheets, restored.worksheets):
            for row in source:
                for cell in row:
                    actual = destination[cell.coordinate]
                    require(actual.value == cell.value and actual.data_type == cell.data_type,
                            "Restored cell value/type differs")
                    checked += 1
        require(restored.properties.creator == NAME, "Metadata restore failed")
        require(restored.active["A2"].comment.author == NAME, "Comment restore failed")
        require(restored.active.oddHeader.center.text == EMAIL, "Header restore failed")
        require(restored.worksheets[1].sheet_state == "veryHidden", "Hidden state lost")
        require(restored.active.row_dimensions[2].hidden, "Hidden row state lost")
        print(json.dumps({"real_model": "PASS", "fixture_categories_checked": 5,
                          "api_round_trip": "PASS", "restored_cells_checked": checked,
                          "residual_scan": "PASS", "offline": True}))
    finally:
        client.delete(url, headers=HEADER)


if __name__ == "__main__":
    main()
