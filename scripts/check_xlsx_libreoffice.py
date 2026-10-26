#!/usr/bin/env python3
"""Synthetic interoperability check using headless LibreOffice, not a model test."""
import io
import json
import subprocess
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src/app"))
from detectors import detect_regex  # noqa: E402
from xlsx_io import WorkbookJob, XlsxError, restore  # noqa: E402
from openpyxl import Workbook, load_workbook  # noqa: E402


def main():
    book = Workbook()
    for cell, value in {"A1": "Email", "A2": "sintetico@example.invalid", "B2": 4,
                        "B3": 5, "C2": "=SUM(B2:B3)", "D1": "Matricola", "D2": 7654321}.items():
        book.active[cell] = value
    raw = io.BytesIO(); book.save(raw)
    job = WorkbookJob(raw.getvalue())
    batch = lambda texts: [detect_regex(t) for t in texts]
    job.analyze(batch, {"columns": [{"sheet": 0, "column": 4}]})
    output = job.export(batch)
    with tempfile.TemporaryDirectory(prefix="rizzo-synthetic-calc-") as temporary:
        root = Path(temporary)
        source = root / "synthetic.xlsx"; source.write_bytes(output)
        destination = root / "converted"; destination.mkdir()
        version = subprocess.run(["libreoffice", "--version"], capture_output=True, text=True, check=True).stdout.strip()
        subprocess.run(["libreoffice", "-env:UserInstallation=" + (root / "profile").as_uri(),
                        "--headless", "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir",
                        str(destination), str(source)], capture_output=True, timeout=60, check=True)
        result = (destination / source.name).read_bytes()
        reopened = load_workbook(io.BytesIO(result), data_only=True)
        original = load_workbook(io.BytesIO(output))
        assert reopened.active["A2"].value == original.active["A2"].value
        assert reopened.active["D2"].value == original.active["D2"].value
        assert reopened.active["C2"].value == 9, "Unmodified formula failed to recalculate"
        try:
            restore(result, job.dictionary)
        except XlsxError:
            pass  # expected: saving in Calc changes the dictionary-bound bytes
        else:
            raise AssertionError("Re-saved workbook was accepted by the original dictionary")
        print(json.dumps({"headless_calc": "PASS", "version": version,
                          "pseudonyms_preserved": True, "formula_recalculation": "PASS",
                          "rewritten_workbook_restore_blocked": True, "detector": "regex-only"}))


if __name__ == "__main__":
    main()
