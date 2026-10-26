#!/usr/bin/env python3
"""Headless local UI smoke test. Optional: pip install playwright.

Start the actual app with local model weights first. Uses a temporary browser
profile and exclusively synthetic data. Does not access a user's browser profile.
All browser requests to other hosts are blocked.
"""
import io
import json
import os
import tempfile
from pathlib import Path
from urllib.parse import urlsplit

from openpyxl import Workbook, load_workbook
from playwright.sync_api import sync_playwright


def main():
    base = os.environ.get("RIZZO_TEST_URL", "http://127.0.0.1:5005").rstrip("/")
    if urlsplit(base).hostname not in {"127.0.0.1", "localhost", "::1"}:
        raise ValueError("The test must target a loopback host")
    book = Workbook()
    for cell, value in {"A1": "Identificativo", "A2": 7654321, "B1": "Email",
                        "B2": "sintetico@example.invalid", "C1": "Nota", "C2": "record-xyz"}.items():
        book.active[cell] = value
    stream = io.BytesIO(); book.save(stream)
    with tempfile.TemporaryDirectory(prefix="rizzo-synthetic-ui-") as temporary, sync_playwright() as p:
        root = Path(temporary)
        browser = p.chromium.launch(executable_path=os.environ.get("RIZZO_CHROME_BIN", "/usr/bin/google-chrome"),
            headless=True, args=["--no-sandbox", "--disable-background-networking", "--disable-sync", "--no-first-run"])
        context = browser.new_context(accept_downloads=True)
        page = context.new_page()
        page.set_default_timeout(150000)
        errors = []
        page.on("pageerror", lambda error: errors.append(str(error)))
        page.route("**/*", lambda route: route.continue_() if urlsplit(route.request.url).netloc ==
                   urlsplit(base).netloc else route.abort())
        page.goto(base + "/xlsx")
        page.locator("#file").set_input_files({"name": "synthetic.xlsx", "mimeType":
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "buffer": stream.getvalue()})
        page.locator("#upload").click()
        page.locator("#configuration").wait_for(state="visible")
        column = page.locator("#sheets fieldset").first.locator(".column").first
        column.locator("input[type=checkbox]").check()
        page.select_option("#language", "en")
        assert page.locator("#sheets .column input[type=checkbox]").first.is_checked(), "Language switch lost configuration"
        page.select_option("#language", "it")
        page.locator("#analyze").click()
        page.locator("#review").wait_for(state="visible")
        page.locator("#working").wait_for(state="hidden")
        assert not page.locator("#downloads").is_visible(), "Unverified downloads visible"
        # Explicit single-occurrence review, then manual addition by coordinates.
        email_row = page.locator("#rows tr").filter(has_text="sintetico@example.invalid").first
        email_row.locator("input:not([type=checkbox])").fill("CONTACT")
        email_row.locator("button").click()
        page.wait_for_function("document.querySelector('#rows input:not([type=checkbox])') !== null")
        page.locator("#addCell").evaluate("e=>e.closest('details').open=true")
        page.locator("#addCell").fill("C2")
        page.locator("#addCategory").fill("NOTE")
        page.locator("#add").click()
        page.wait_for_function("document.querySelector('#rows').textContent.includes('NOTE')")
        page.locator("#export").click()
        page.locator("#downloads").wait_for(state="visible")
        page.locator("#working").wait_for(state="hidden")
        with page.expect_download() as event:
            page.locator("#download").click()
        event.value.save_as(root / "masked.xlsx")
        with page.expect_download() as event:
            page.locator("#downloadMapping").click()
        event.value.save_as(root / "dictionary.json")
        output = load_workbook(root / "masked.xlsx")
        assert output.active["A2"].value != 7654321
        assert output.active["B2"].value.startswith("[[RZ_")
        assert output.active["C2"].value.startswith("[[RZ_")
        page.locator("#restoreFile").set_input_files(root / "masked.xlsx")
        page.locator("#restoreMapping").set_input_files(root / "dictionary.json")
        with page.expect_download() as event:
            page.locator("#restore").click()
        event.value.save_as(root / "restored.xlsx")
        restored = load_workbook(root / "restored.xlsx")
        for row in book.active:
            for cell in row:
                actual = restored.active[cell.coordinate]
                assert actual.value == cell.value and actual.data_type == cell.data_type
        page.locator("#remove").click()
        page.locator("#configuration").wait_for(state="hidden")
        assert not errors, "Browser JavaScript errors"
        browser.close()
        print(json.dumps({"browser_ui": "PASS", "real_server": base, "round_trip_cells": 6,
                          "javascript_errors": len(errors), "external_browser_requests_allowed": False}))


if __name__ == "__main__":
    main()
